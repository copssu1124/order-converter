# -*- coding: utf-8 -*-
"""ERP 등록 · 원장 비교 · 제이 품목거래 — 직원이 만든 '제이제이 오프라인 업무도구 v3.7.2'(PowerShell+Excel COM)를
   엑셀 설치 없이 돌아가도록 파이썬으로 옮긴 것. 규칙은 그 도구의 사용안내.txt/JJ_Offline_Tool.ps1 와 동일하게 유지.

입력: 우리 변환기의 매입_*.xls · 매출_*.xls, 원본 주문서(.xlsx), 원장(.xls, 다모아/제이제이/제이무역 선택), 품목정보관리.xls(ERP 단가표)
출력: ERP_판매구매일괄등록_날짜_시각.xls(97-2003 양식) · ERP_확인필요_….xlsx · 원장비교결과_….xlsx · 제이 품목거래 월-일.xls
"""
import os
import re
import math
import datetime
import xlrd
import xlwt
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from xlutils.copy import copy as xl_copy

# ── 거래처 코드 (ERP 입력 규칙) ──
PARTNER_JJ = ('00787', '청년몰-제이')
PARTNER_DAMOA = ('00789', '청년몰-다모아')
PARTNER_TRADE = ('01147', '청년몰-제이무역')
PARTNER_COUPANG = ('01082', '쿠팡')
PARTNER_STORE = ('00326', '스마트 스토어')
# 서비스 품목 코드
CODE_SETTLE_STORE, CODE_SETTLE_COUPANG = '0000000128', '0000000581'
CODE_VEHICLE, CODE_DAMOA_SHIP, CODE_JJ_SHIP, CODE_JJ_3PL, CODE_CANOPY = '0000000567', '0000001714', '0000001715', '0000001716', '0000000723'
CODE_JJ_3PL_FREIGHT = '0000000602'      # JJ 3PL/화물 — 용차(운임/1) 거래의 청년몰-제이 3PL 줄 (직원 요청 2026-10-02)

OWNER_COLORS = {'다모아': 'DDEBF7', '제이제이': 'E2EFDA', '제이무역': 'FFF2CC', '미분류': 'F2F2F2'}
SERVICE_COLORS = {'다모아': '9DC3E6', '제이제이': 'A9D08E', '제이무역': 'F4B183'}
MISMATCH_COLOR, HILITE_COLOR, HEADER_COLOR = 'F4CCCC', 'FFFFCC', '305496'
PURCHASE_SUMMARY_COLOR, SALES_SUMMARY_COLOR = 'DDEBF7', 'FCE4EC'
_ILLEGAL = re.compile('[' + ''.join(chr(c) for c in range(32) if c not in (9, 10, 13)) + ']')   # 원장(.xls) 셀의 제어문자 제거(openpyxl 저장 오류 방지)


# ═══════════ 공통 유틸 ═══════════
def _num(v):
    if v is None:
        return 0.0
    s = str(v).replace(',', '').strip()
    if s == '':
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _round(x):
    """Excel/.NET AwayFromZero 반올림 (파이썬 round는 짝수 반올림이라 다름)."""
    x = float(x)
    return float(math.floor(abs(x) + 0.5)) * (1 if x >= 0 else -1)


def _near(a, b, tol=0.01):
    return abs(float(a) - float(b)) <= tol


def _rows(path, sheet=None):
    """(.xls/.xlsx 공통) 시트 → 1-based 접근 가능한 2차원 리스트. 반환: {name: rows}. rows[r-1][c-1]."""
    out = {}
    if path.lower().endswith('.xls'):
        wb = xlrd.open_workbook(path)

        def val(sh, r, c):
            ct, v = sh.cell_type(r, c), sh.cell_value(r, c)
            if ct == xlrd.XL_CELL_DATE:
                try:
                    return xlrd.xldate_as_datetime(v, wb.datemode)
                except Exception:
                    return v
            if isinstance(v, str):
                return _ILLEGAL.sub('', v)
            return v
        for sh in wb.sheets():
            if sheet and sh.name != sheet:
                continue
            out[sh.name] = [[val(sh, r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if sheet and ws.title != sheet:
                continue
            out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    return out


def _cell(rows, r, c):
    try:
        return rows[r - 1][c - 1]
    except IndexError:
        return None


def trade_date_from_name(path):
    """주문서 파일명의 M.D / YYYY.M.D 에서 거래일자 추출 (없으면 None)."""
    name = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r'(?<!\d)(20\d{2})[.\-_](\d{1,2})[.\-_](\d{1,2})(?!\d)', name)
    try:
        if m:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = re.search(r'(?<!\d)(\d{1,2})[.\-_](\d{1,2})(?!\d)', name)
        if m:
            return datetime.date(datetime.date.today().year, int(m.group(1)), int(m.group(2)))
    except ValueError:
        pass
    return None


# ═══════════ 품목명 정규화 (도구와 동일 규칙) ═══════════
def owner_item_name(name, owner):
    v = _s(name)
    if owner == '다모아' and '완제품' in v and not re.match(r'^\s*DA/', v):
        return 'DA/' + v
    return v


def canonical(name, owner=''):
    v = owner_item_name(name, owner).strip().lower().replace('（', '(').replace('）', ')')
    v = re.sub(r'\(이면은박\)$', '', v)
    v = re.sub(r'\(ds-1\)', '', v)
    v = re.sub(r'\(김치15k\)', '', v)
    v = re.sub(r'-김치15k$', '', v)
    v = re.sub(r'\s+', '', v)
    v = re.sub(r'/hs$', '', v)
    if v == '3k-t5/ts-27':
        v = '3k-t5/ts'
    return v


def is_excluded_item(name):
    v = _s(name)
    return v == '' or re.search(r'택배비|3\s*pl|천막추가옵션|파손/재발송', v, re.I) is not None


# ═══════════ 적요 · 거래처 · 시트명 매핑 ═══════════
def order_sheet_name(alias):
    """주문서 B열(별칭) → 분리 시트명 (도구의 Resolve-OrderSheetName)."""
    b = _s(alias)
    rules = [
        (r'^다모아스토어$', lambda m: '다모아스토어'), (r'^다모아쿠팡$', lambda m: '다모아쿠팡'),
        (r'^다모아무역스토어$', lambda m: '제이무역스토어-다모아'), (r'^다모아무역쿠팡$', lambda m: '제이무역쿠팡-다모아'),
        (r'^다모아개인/(.+)$', lambda m: '다모아개인 ' + m.group(1)),
        (r'^제이스토어$', lambda m: '제이스토어'), (r'^제이쿠팡$', lambda m: '제이쿠팡'),
        (r'^제이무역스토어$', lambda m: '제이무역스토어'), (r'^제이무역쿠팡$', lambda m: '제이무역쿠팡'),
        (r'^제이무역스토어[/-](.+)$', lambda m: '제이무역스토어-' + m.group(1)),
        (r'^제이무역쿠팡[/-](.+)$', lambda m: '제이무역쿠팡-' + m.group(1)),
        (r'^무역\s*스토어$', lambda m: '제이무역스토어'), (r'^무역\s*쿠팡$', lambda m: '제이무역쿠팡'),
        (r'^제이개인/(.+)$', lambda m: '제이개인 ' + m.group(1)),
        (r'^용차스토어/다모아/(.+)$', lambda m: '다모아스토어-' + m.group(1)),
        (r'^용차스토어/제이/(.+)$', lambda m: '제이스토어-' + m.group(1)),
        (r'^용차쿠팡/다모아/(.+)$', lambda m: '다모아쿠팡-' + m.group(1)),
        (r'^용차쿠팡/제이/(.+)$', lambda m: '제이쿠팡-' + m.group(1)),
        (r'^용차개인/다모아/(.+)$', lambda m: '다모아개인 ' + m.group(1)),
        (r'^용차개인/제이/(.+)$', lambda m: '제이개인 ' + m.group(1)),
    ]
    for pat, fn in rules:
        m = re.match(pat, b)
        if m:
            return fn(m)
    return '미분류'


def normalize_remark(name):
    v = _s(name)
    m = re.match(r'^청년몰-제이무역\s+(스토어|쿠팡)-다모아$', v)
    if m:
        return '청년몰-제이무역 %s/다모아' % m.group(1)
    m = re.match(r'^청년몰-제이무역\s+(스토어|쿠팡)-(.+)$', v)
    if m:
        return '청년몰-제이무역 %s/%s' % (m.group(1), m.group(2))
    if v.startswith('청년몰-'):
        m = re.match(r'^(청년몰-(?:다모아|제이)\s+(?:스토어|쿠팡|개인))-(.+)$', v)
        return m.group(1) + '/' + m.group(2) if m else v
    rules = [
        (r'^다모아무역스토어$', lambda m: '청년몰-제이무역 스토어/다모아'), (r'^다모아무역쿠팡$', lambda m: '청년몰-제이무역 쿠팡/다모아'),
        (r'^다모아스토어$', lambda m: '청년몰-다모아 스토어'), (r'^다모아쿠팡$', lambda m: '청년몰-다모아 쿠팡'),
        (r'^제이스토어$', lambda m: '청년몰-제이 스토어'), (r'^제이쿠팡$', lambda m: '청년몰-제이 쿠팡'),
        (r'^제이무역스토어-다모아$', lambda m: '청년몰-제이무역 스토어/다모아'), (r'^제이무역쿠팡-다모아$', lambda m: '청년몰-제이무역 쿠팡/다모아'),
        (r'^제이무역스토어$', lambda m: '청년몰-제이무역 스토어'), (r'^제이무역쿠팡$', lambda m: '청년몰-제이무역 쿠팡'),
        (r'^제이무역스토어-(.+)$', lambda m: '청년몰-제이무역 스토어/' + m.group(1)),
        (r'^제이무역쿠팡-(.+)$', lambda m: '청년몰-제이무역 쿠팡/' + m.group(1)),
        (r'^무역\s*스토어$', lambda m: '청년몰-제이무역 스토어'), (r'^무역\s*쿠팡$', lambda m: '청년몰-제이무역 쿠팡'),
        (r'^다모아개인[\s\-]+(.+)$', lambda m: '청년몰-다모아 개인/' + m.group(1)),
        (r'^제이개인[\s\-]+(.+)$', lambda m: '청년몰-제이 개인/' + m.group(1)),
        (r'^다모아스토어-(.+)$', lambda m: '청년몰-다모아 스토어/' + m.group(1)),
        (r'^다모아쿠팡-(.+)$', lambda m: '청년몰-다모아 쿠팡/' + m.group(1)),
        (r'^제이스토어-(.+)$', lambda m: '청년몰-제이 스토어/' + m.group(1)),
        (r'^제이쿠팡-(.+)$', lambda m: '청년몰-제이 쿠팡/' + m.group(1)),
    ]
    for pat, fn in rules:
        m = re.match(pat, v)
        if m:
            return fn(m)
    return v


def purchase_sheet_to_remark(sheet_name):
    """매입파일 시트명 → ERP 적요. 변환 불가면 ValueError."""
    n = _s(sheet_name)
    if n.startswith('청년몰-'):
        return normalize_remark(n)
    rules = [
        (r'^무역\s*스토어$', lambda m: '청년몰-제이무역 스토어/다모아'), (r'^무역\s*쿠팡$', lambda m: '청년몰-제이무역 쿠팡/다모아'),
        (r'^다모아무역스토어$', lambda m: '청년몰-제이무역 스토어/다모아'), (r'^다모아무역쿠팡$', lambda m: '청년몰-제이무역 쿠팡/다모아'),
        (r'^제이무역스토어[-/]다모아$', lambda m: '청년몰-제이무역 스토어/다모아'), (r'^제이무역쿠팡[-/]다모아$', lambda m: '청년몰-제이무역 쿠팡/다모아'),
        (r'^제이무역스토어[-/](.+)$', lambda m: '청년몰-제이무역 스토어/' + m.group(1)),
        (r'^제이무역쿠팡[-/](.+)$', lambda m: '청년몰-제이무역 쿠팡/' + m.group(1)),
        (r'^제이무역스토어$', lambda m: '청년몰-제이무역 스토어'), (r'^제이무역쿠팡$', lambda m: '청년몰-제이무역 쿠팡'),
        (r'^다모아스토어$', lambda m: '청년몰-다모아 스토어'), (r'^다모아쿠팡$', lambda m: '청년몰-다모아 쿠팡'),
        (r'^제이스토어$', lambda m: '청년몰-제이 스토어'), (r'^제이쿠팡$', lambda m: '청년몰-제이 쿠팡'),
        (r'^다모아개인[\s/\-]+(.+)$', lambda m: '청년몰-다모아 개인/' + m.group(1)),
        (r'^제이개인[\s/\-]+(.+)$', lambda m: '청년몰-제이 개인/' + m.group(1)),
        (r'^용차스토어[\s/\-]+다모아[\s/\-]+(.+)$', lambda m: '청년몰-다모아 스토어/' + m.group(1)),
        (r'^용차스토어[\s/\-]+제이[\s/\-]+(.+)$', lambda m: '청년몰-제이 스토어/' + m.group(1)),
        (r'^용차쿠팡[\s/\-]+다모아[\s/\-]+(.+)$', lambda m: '청년몰-다모아 쿠팡/' + m.group(1)),
        (r'^용차쿠팡[\s/\-]+제이[\s/\-]+(.+)$', lambda m: '청년몰-제이 쿠팡/' + m.group(1)),
        (r'^용차개인[\s/\-]+다모아[\s/\-]+(.+)$', lambda m: '청년몰-다모아 개인/' + m.group(1)),
        (r'^용차개인[\s/\-]+제이[\s/\-]+(.+)$', lambda m: '청년몰-제이 개인/' + m.group(1)),
    ]
    for pat, fn in rules:
        m = re.match(pat, n)
        if m:
            return fn(m)
    raise ValueError('매입파일 시트명을 ERP 적요로 변환할 수 없습니다: ' + n)


def owner_from_remark(remark):
    if remark.startswith('청년몰-제이무역'):
        return '제이무역'
    if remark.startswith('청년몰-다모아'):
        return '다모아'
    return '제이제이'


def purchase_partner(remark):
    if remark.startswith('청년몰-제이무역'):
        return PARTNER_TRADE
    if remark.startswith('청년몰-다모아'):
        return PARTNER_DAMOA
    if remark.startswith('청년몰-제이'):
        return PARTNER_JJ
    return None


def sales_partner(remark):
    if '개인' in remark:
        return None
    if '쿠팡' in remark:
        return PARTNER_COUPANG
    if '스토어' in remark:
        return PARTNER_STORE
    return None


def service_brand(remark):
    return '다모아' if (remark.startswith('청년몰-다모아') or re.search(r'(?:-|/)다모아$', remark)) else '제이'


def purchase_owner(sheet_name):
    n = _s(sheet_name)
    if '무역' in n:
        return '제이무역'
    if re.match(r'^다모아', n) or re.search(r'다모아\s', n):
        return '다모아'
    if re.match(r'^제이', n) or re.search(r'제이\s', n):
        return '제이제이'
    return '미분류'


def shipping_ledger_owner(sheet_name, product_owner):
    if '용차' in sheet_name:
        return '제외'
    if '무역' in sheet_name:
        return '다모아'
    if product_owner in ('다모아', '제이제이'):
        return product_owner
    return '제이무역'


def order_shipping_category(alias):
    v = _s(alias)
    rules = [
        (r'^다모아무역스토어$', ('다모아', '제이무역스토어/다모아')), (r'^다모아무역쿠팡$', ('다모아', '제이무역쿠팡/다모아')),
        (r'^제이무역스토어(?:[/-].+)?$', ('제이무역', '제이무역스토어')), (r'^제이무역쿠팡(?:[/-].+)?$', ('제이무역', '제이무역쿠팡')),
        (r'^무역\s*스토어$', ('다모아', '제이무역스토어/다모아')), (r'^무역\s*쿠팡$', ('다모아', '제이무역쿠팡/다모아')),
        (r'^제이스토어(?:[/-].+)?$', ('제이제이', '제이스토어')), (r'^제이쿠팡(?:[/-].+)?$', ('제이제이', '제이쿠팡')),
        (r'^제이개인(?:[/-].+)?$', ('제이제이', '제이개인')),
        (r'^다모아스토어(?:[/-].+)?$', ('다모아', '다모아스토어')), (r'^다모아쿠팡(?:[/-].+)?$', ('다모아', '다모아쿠팡')),
        (r'^다모아개인(?:[/-].+)?$', ('다모아', '다모아개인')),
        (r'^용차스토어/다모아(?:/.+)?$', ('다모아', '다모아스토어')), (r'^용차쿠팡/다모아(?:/.+)?$', ('다모아', '다모아쿠팡')),
        (r'^용차개인/다모아(?:/.+)?$', ('다모아', '다모아개인')),
        (r'^용차스토어/제이(?:/.+)?$', ('제이제이', '제이스토어')), (r'^용차쿠팡/제이(?:/.+)?$', ('제이제이', '제이쿠팡')),
        (r'^용차개인/제이(?:/.+)?$', ('제이제이', '제이개인')),
    ]
    for pat, res in rules:
        if re.match(pat, v):
            return res
    return ('미분류', v or '미분류')


# ═══════════ ERP 단가표 ═══════════
def load_erp_map(path):
    """품목정보관리.xls → {canonical: [item, …]}. 데이터는 7행부터(A코드 B품목 C규격 D단위 G종류 H과세 I바코드 J구매단가 K판매단가)."""
    rows = next(iter(_rows(path).values()))
    erp = {}
    for r in range(7, len(rows) + 1):
        name = _s(_cell(rows, r, 2))
        if name == '' or re.search(r'/s$', name):
            continue
        price = _num(_cell(rows, r, 10))
        item = {'code': _s(_cell(rows, r, 1)), 'name': name, 'detail': _s(_cell(rows, r, 3)), 'unit': _s(_cell(rows, r, 4)),
                'kind': _s(_cell(rows, r, 7)), 'barcode': _s(_cell(rows, r, 9)), 'purchase': price,
                'sale': _num(_cell(rows, r, 11)), 'rounded': _round(price)}
        erp.setdefault(canonical(name), []).append(item)
    return erp


def _candidates(erp, key):
    allc = erp.get(key) or []
    prods = [x for x in allc if x['kind'] == '제품']
    return prods if prods else allc


def resolve_erp_item(erp, key):
    """원장 비교용 (구매단가 기준)."""
    if key not in erp:
        return {'status': 'ERP 미등록', 'item': None, 'price': None, 'note': ''}
    c = _candidates(erp, key)
    prices = sorted({x['rounded'] for x in c})
    if len(prices) != 1:
        return {'status': 'ERP 단가 중복', 'item': c[0], 'price': None, 'note': ', '.join(_s(p) for p in prices)}
    return {'status': '정상', 'item': c[0], 'price': prices[0], 'note': ''}


def resolve_export_item(erp, key, mode):
    if key not in erp:
        return {'status': 'ERP 미등록', 'item': None, 'price': None}
    c = _candidates(erp, key)
    if mode == '판매0원':
        return {'status': '정상', 'item': c[0], 'price': 0.0}
    field = 'purchase' if mode == '구매' else 'sale'
    prices = sorted({_round(x[field]) for x in c})
    if len(prices) != 1:
        return {'status': 'ERP 단가 중복', 'item': c[0], 'price': None}
    return {'status': '정상', 'item': c[0], 'price': prices[0]}


# ═══════════ 거래 읽기 ═══════════
def _add_grouped(index, items, name, qty, owner):
    if qty == 0 or is_excluded_item(name):
        return
    conv = owner_item_name(name, owner)
    key = canonical(conv, owner)
    if not key:
        return
    if key in index:
        index[key]['qty'] += qty
    else:
        it = {'key': key, 'name': conv, 'qty': qty, 'erp': None}
        index[key] = it
        items.append(it)


def _service(items, name, code, qty, unit_price=None, supply=None, unit=None, fixed=None):
    items.append({'key': canonical(name), 'name': name, 'qty': qty, 'erp': None, 'use_purchase': True,
                  'expected': code, 'unit_price': unit_price, 'supply': supply, 'unit': unit, 'fixed': fixed})


def read_purchase_transactions(path):
    out = []
    for sheet, rows in _rows(path).items():
        is_vehicle = '용차' in sheet
        remark = purchase_sheet_to_remark(sheet)
        partner = purchase_partner(remark)
        if partner is None:
            raise ValueError('구매 거래처를 결정할 수 없습니다: ' + remark)
        owner = owner_from_remark(remark)
        items, index, ship, grade = [], {}, 0.0, 0.0
        for r in range(1, len(rows) + 1):
            name = _s(_cell(rows, r, 4)); qty = _num(_cell(rows, r, 7))
            amount = _num(_cell(rows, r, 18)); label = _s(_cell(rows, r, 19))
            if '배송비합계' in label:
                ship += amount
            elif '택배등급' in label:
                grade += amount
            _add_grouped(index, items, name, qty, owner)
        if items:
            out.append({'mode': '구매', 'remark': remark, 'partner': partner, 'owner': owner, 'items': items,
                        'shipping': ship, 'grade': grade, 'service_total': ship + grade, 'vehicle': is_vehicle, 'vehicle_shipping': 0.0})
    return out


def parse_order_item(text, fallback_qty):
    v = _s(text)
    if is_excluded_item(v):
        return None
    mult = _num(fallback_qty)
    tail = re.search(r'\*(\d+(?:\.\d+)?)\s*$', v)
    if tail:
        mult = float(tail.group(1))
    m = re.search(r'\((?:\d+(?:\.\d+)?)\s*박스\s*/\s*(\d+(?:\.\d+)?)\s*매\)\s*\*\d+(?:\.\d+)?\s*$', v, re.I)
    if m:
        return v[:m.start()].strip(), float(m.group(1)) * mult
    m = re.search(r'\((\d+(?:\.\d+)?)\s*(?:ea|매)\)\s*\*\d+(?:\.\d+)?\s*$', v, re.I)
    if m:
        return v[:m.start()].strip(), float(m.group(1)) * mult
    if tail:
        return v[:tail.start()].strip(), mult
    return v, mult


def _order_rows(order_path):
    sheets = _rows(order_path)
    for nm in ('주문서 붙여넣기', '주문서'):
        if nm in sheets:
            return sheets[nm]
    return next(iter(sheets.values()))


def read_sales_transactions(order_path):
    """원본 주문서를 B열로 묶어 판매 거래로. (도구의 '시트분리' 중간 파일 없이 바로 같은 결과)"""
    rows = _order_rows(order_path)
    groups = {}
    for r in range(2, len(rows) + 1):
        b = _s(_cell(rows, r, 2))
        if b == '':
            continue
        groups.setdefault(order_sheet_name(b), []).append(r)
    trans, qty_by, veh_by, ship_by = [], {}, {}, {}
    for gname in sorted(groups):
        remark = normalize_remark(gname)
        partner = sales_partner(remark)
        owner = owner_from_remark(remark)
        items, index, order_count, is_vehicle, ship_total = [], {}, 0.0, False, 0.0
        for r in groups[gname]:
            origin, alias = _s(_cell(rows, r, 1)), _s(_cell(rows, r, 2))
            if '용차' in origin or '용차' in alias:
                is_vehicle = True
            if alias:
                ship_total += _num(_cell(rows, r, 18))
            parsed = parse_order_item(_cell(rows, r, 9), _cell(rows, r, 11))
            if parsed:
                order_count += _num(_cell(rows, r, 11))
                if partner:
                    _add_grouped(index, items, parsed[0], parsed[1], owner)
        qty_by[remark] = qty_by.get(remark, 0.0) + order_count
        ship_by[remark] = ship_by.get(remark, 0.0) + ship_total
        if is_vehicle:
            veh_by[remark] = True
        if partner and items:
            trans.append({'mode': '판매', 'remark': remark, 'partner': partner, 'owner': owner, 'items': items,
                          'vehicle': is_vehicle, 'vehicle_shipping': ship_total, 'service_total': 0.0})
    return trans, qty_by, veh_by, ship_by


def read_settlement_totals(sales_path):
    totals = {}
    for sheet, rows in _rows(sales_path).items():
        remark = normalize_remark(sheet)
        if '개인' in remark:
            continue
        if not (re.match(r'^청년몰-(?:제이|다모아) (?:스토어|쿠팡)$', remark) or re.match(r'^청년몰-제이무역 (?:스토어|쿠팡)(?:/.+)?$', remark)):
            continue
        total, found = 0.0, False
        for r in range(1, len(rows) + 1):
            label = re.sub(r'\s', '', _s(_cell(rows, r, 19)))
            if re.search(r'배송비.*정산예정금액.*토탈', label) or re.match(r'^배송비\+?정산예정금액토탈$', label):
                total += _num(_cell(rows, r, 18)); found = True
        if found:
            totals[remark] = totals.get(remark, 0.0) + total
    return totals


def _add_service_rows(t, service_total, order_count):
    brand = service_brand(t['remark'])
    if t.get('vehicle'):
        _service(t['items'], '운/1', CODE_VEHICLE, t.get('vehicle_shipping', 0.0), 0.0, unit='대')
        if brand != '다모아' and t['remark'].startswith('청년몰-제이') and not t['remark'].startswith('청년몰-제이무역'):
            # 용차 + 청년몰-제이 → 택배용 '제이 3PL' 대신 화물용 (구매·판매 공통. 판매는 거래처가 스토어/쿠팡이라 적요로 판단)
            _service(t['items'], 'JJ 3PL/화물', CODE_JJ_3PL_FREIGHT, order_count)
        elif brand != '다모아':
            _service(t['items'], '제이 3PL', CODE_JJ_3PL, order_count)
        return
    if brand == '다모아':
        if re.match(r'^청년몰-제이무역\s+(스토어|쿠팡)/다모아$', t['remark']):
            _service(t['items'], '다모아 택배&외박스', CODE_DAMOA_SHIP, service_total, 1.0)   # 구매 단가 1 (직원 요청 2026-10-02, 판매는 아래에서 항상 0)
        elif re.match(r'^청년몰-다모아\s+개인/', t['remark']):
            _service(t['items'], '다모아 택배&외박스', CODE_DAMOA_SHIP, None, 1.0, supply=service_total)
        else:
            _service(t['items'], '다모아 택배&외박스', CODE_DAMOA_SHIP, service_total, 1.0)
    else:
        _service(t['items'], '제이 택배', CODE_JJ_SHIP, service_total, 1.0)
        _service(t['items'], '제이 3PL', CODE_JJ_3PL, order_count)


def _add_settlement(t, amount):
    if t['mode'] != '판매' or '개인' in t['remark']:
        return
    rk = t['remark']
    if re.match(r'^청년몰-(?:제이|다모아) 스토어$', rk) or re.match(r'^청년몰-제이무역 스토어(?:/.+)?$', rk):
        fixed = {'code': CODE_SETTLE_STORE, 'barcode': '', 'name': '스토어정산예정금액', 'detail': '', 'unit': '건'}
        _service(t['items'], '스토어정산예정금액', CODE_SETTLE_STORE, 1.0, amount, unit='건', fixed=fixed)
    elif re.match(r'^청년몰-(?:제이|다모아) 쿠팡$', rk) or re.match(r'^청년몰-제이무역 쿠팡(?:/.+)?$', rk):
        fixed = {'code': CODE_SETTLE_COUPANG, 'barcode': '', 'name': '쿠팡정산', 'detail': '', 'unit': '원'}
        _service(t['items'], '쿠팡정산', CODE_SETTLE_COUPANG, 1.0, amount, unit='원', fixed=fixed)


# ═══════════ ① ERP 판매·구매 일괄등록 파일 ═══════════
def _xls_style(bold=False, color=None, numfmt=None, textfmt=False, datefmt=False):
    st = xlwt.XFStyle()
    if bold:
        f = xlwt.Font(); f.bold = True; st.font = f
    if color:
        p = xlwt.Pattern(); p.pattern = xlwt.Pattern.SOLID_PATTERN; p.pattern_fore_colour = color; st.pattern = p
    if numfmt:
        st.num_format_str = numfmt
    if textfmt:
        st.num_format_str = '@'
    if datefmt:
        st.num_format_str = 'yyyy-mm-dd'
    return st


def _xls_colors(wb):
    """xlwt 팔레트에 거래처 색을 등록하고 색 인덱스를 돌려줌."""
    idx = {}
    n = 8   # 8~63 사용자 색 자리
    for name, hexc in list(OWNER_COLORS.items()) + [('svc_' + k, v) for k, v in SERVICE_COLORS.items()]:
        r, g, b = int(hexc[0:2], 16), int(hexc[2:4], 16), int(hexc[4:6], 16)
        xlwt.add_palette_colour('c_%s' % name, n)
        wb.set_colour_RGB(n, r, g, b)
        idx[name] = n
        n += 1
    return idx


def build_erp_upload(purchase_path, sales_path, order_path, erp_path, template_path, trade_date, out_dir=None, log=print):
    """ERP_판매구매일괄등록_날짜_시각.xls (+ ERP_확인필요_….xlsx) 생성. 반환 dict(out, issues_path, issues, 거래수, 품목수, 미리보기행)."""
    for p in (purchase_path, sales_path, order_path, erp_path, template_path):
        if not os.path.exists(p):
            raise FileNotFoundError('필요한 파일을 찾을 수 없습니다: ' + p)
    out_dir = out_dir or os.path.dirname(purchase_path)
    stamp = datetime.datetime.now().strftime('%H%M%S')
    erp = load_erp_map(erp_path)
    log('ERP 품목 %d개 키 로드' % len(erp))
    purchases = read_purchase_transactions(purchase_path)
    sales, qty_by, veh_by, ship_by = read_sales_transactions(order_path)
    for t in purchases:
        if t['remark'] in veh_by:
            t['vehicle'] = True
            t['vehicle_shipping'] = ship_by.get(t['remark'], 0.0)
    settle = read_settlement_totals(sales_path)
    issues = []
    purchase_by = {}
    for t in purchases:
        if t['remark'] not in purchase_by or not t['vehicle']:
            purchase_by[t['remark']] = t
    for t in purchases:
        _add_service_rows(t, t['service_total'], qty_by.get(t['remark'], 0.0))
    for t in sales:
        st = purchase_by[t['remark']]['service_total'] if t['remark'] in purchase_by else 0.0
        _add_service_rows(t, st, qty_by.get(t['remark'], 0.0))
        rk = t['remark']
        if re.match(r'^청년몰-(?:제이|다모아) (?:스토어|쿠팡)$', rk) or re.match(r'^청년몰-제이무역 (?:스토어|쿠팡)(?:/.+)?$', rk):
            if rk in settle:
                _add_settlement(t, settle[rk])
            else:
                _add_settlement(t, 0.0)
                issues.append({'mode': '판매', 'remark': rk, 'code': t['partner'][0], 'partner': t['partner'][1],
                               'item': '쿠팡정산' if '쿠팡' in rk else '스토어정산예정금액', 'qty': 1,
                               'error': '매출파일에서 배송비+정산예정금액 토탈을 찾지 못함', 'key': '매출 정산금액'})
    # 천막 택배 (제이무역/다모아 조합이 있을 때)
    canopy = [t for t in purchases if not t['vehicle'] and re.match(r'^청년몰-제이무역\s+(스토어|쿠팡)/다모아$', t['remark'])]
    if canopy:
        total = sum(t['service_total'] for t in canopy)
        items = []
        _service(items, '천막 택배', CODE_CANOPY, total, 1.0)
        purchases.append({'mode': '구매', 'remark': '청년몰-다모아 천막택배비*', 'partner': PARTNER_DAMOA, 'owner': '다모아',
                          'items': items, 'service_total': total, 'vehicle': False, 'vehicle_shipping': 0.0})
    trans = purchases + sales
    if not trans:
        raise ValueError('ERP 입력자료로 변환할 거래가 없습니다.')
    trans.sort(key=lambda t: (t['remark'], 0 if t['mode'] == '구매' else 1))
    # 품목코드·단가 확인
    for t in trans:
        for it in t['items']:
            if it.get('fixed'):
                it['erp'] = {'status': '정상', 'item': it['fixed'], 'price': float(it['unit_price'] or 0)}
                continue
            zero_sale = (t['mode'] == '판매')
            mode = '판매0원' if zero_sale else ('구매' if it.get('use_purchase') else t['mode'])
            res = resolve_export_item(erp, it['key'], mode)
            if res['status'] != '정상':
                issues.append({'mode': t['mode'], 'remark': t['remark'], 'code': t['partner'][0], 'partner': t['partner'][1],
                               'item': it['name'], 'qty': it['qty'], 'error': res['status'], 'key': it['key']})
                continue
            if it.get('expected') and res['item']['code'] != it['expected']:
                issues.append({'mode': t['mode'], 'remark': t['remark'], 'code': t['partner'][0], 'partner': t['partner'][1],
                               'item': it['name'], 'qty': it['qty'],
                               'error': '품목코드 불일치 (%s, 기대값 %s)' % (res['item']['code'], it['expected']), 'key': it['key']})
                continue
            it['erp'] = res
            if not zero_sale and it.get('unit_price') is not None:
                it['erp']['price'] = float(it['unit_price'])
    valid_lines = sum(1 for t in trans for it in t['items'] if it['erp'])
    valid_trans = sum(1 for t in trans if any(it['erp'] for it in t['items']))
    issues_path = None
    if issues:
        issues_path = os.path.join(out_dir, 'ERP_확인필요_%s_%s.xlsx' % (trade_date.strftime('%Y%m%d'), stamp))
        _write_issues(issues, issues_path)
        log('ERP 확인필요 %d건: %s' % (len(issues), os.path.basename(issues_path)))
    if valid_lines == 0:
        return {'out': None, 'issues_path': issues_path, 'issues': issues, 'transactions': 0, 'lines': 0, 'rows': []}

    # 양식 복사(서식·작성예제 유지) 후 3행부터 기록
    out_path = os.path.join(out_dir, 'ERP_판매구매일괄등록_%s_%s.xls' % (trade_date.strftime('%Y%m%d'), stamp))
    rb = xlrd.open_workbook(template_path, formatting_info=True)
    wb = xl_copy(rb)
    ws = wb.get_sheet(rb.sheet_names().index('대량등록양식'))
    colors = _xls_colors(wb)
    sty = {'txt': _xls_style(textfmt=True), 'date': _xls_style(datefmt=True), 'num': _xls_style(numfmt='#,##0.###'), 'plain': _xls_style()}
    row, preview = 3, []
    for t in trans:
        valid = [it for it in t['items'] if it['erp']]
        if not valid:
            continue
        first = True
        for it in valid:
            info, unit_price = it['erp']['item'], float(it['erp']['price'] or 0)
            qty = None if it['qty'] is None else float(it['qty'])
            sup_over = it.get('supply')
            is_settle = bool(it.get('fixed'))
            if t['mode'] == '판매' and not is_settle:
                unit_price, supply, vat = 0.0, 0.0, 0.0
            elif t['mode'] == '구매':
                supply = float(sup_over) if sup_over is not None else _round((qty or 0) * unit_price)
                vat = _round(supply * 0.1)
            else:
                if sup_over is not None:
                    supply = float(sup_over); vat = _round(supply * 0.1)
                else:
                    total = _round((qty or 0) * unit_price)
                    supply = _round(total / 1.1); vat = total - supply
            is_service = re.match(r'^(다모아\s*택배&외박스|제이\s*택배|제이\s*3\s*PL|천막\s*택배)$', info['name']) is not None
            color = None
            if is_service and t['owner'] in SERVICE_COLORS:
                color = colors['svc_' + t['owner']]
            elif t['mode'] == '구매' and t['owner'] in OWNER_COLORS:
                color = colors[t['owner']]
            def put(c, v, kind='plain'):
                st = _xls_style(bold=is_service, color=color, textfmt=(kind == 'txt'), datefmt=(kind == 'date'),
                                numfmt=('#,##0.###' if kind == 'num' else None))
                ws.write(row - 1, c - 1, v, st)
            for c in range(1, 26):
                put(c, '')
            if first:
                put(1, t['mode']); put(2, t['partner'][0], 'txt'); put(3, t['partner'][1])
                put(4, datetime.datetime(trade_date.year, trade_date.month, trade_date.day), 'date')
                put(5, '상품매출' if t['mode'] == '판매' else '상품')
                put(6, '건별/부가세포함' if t['mode'] == '판매' else '과세/부가세별도10%')
                put(7, t['remark']); put(13, '본사')
            put(14, info['code'], 'txt'); put(15, info['barcode'], 'txt'); put(16, info['name']); put(17, info['detail'])
            put(18, it.get('unit') or info['unit'])
            if qty is not None:
                put(19, qty, 'num')
            put(22, unit_price, 'num'); put(23, supply, 'num'); put(24, vat, 'num')
            preview.append((t['mode'] if first else '', t['remark'] if first else '', info['name'], qty, unit_price, supply, vat))
            first = False; row += 1
    for c in range(25):
        ws.col(c).width = 256 * 14
    ws.col(15).width = 256 * 30
    wb.save(out_path)
    log('ERP 입력파일 완료: %s (정상 거래 %d건 · 품목 %d건)' % (os.path.basename(out_path), valid_trans, valid_lines))
    return {'out': out_path, 'issues_path': issues_path, 'issues': issues, 'transactions': valid_trans, 'lines': valid_lines, 'rows': preview}


def _write_issues(issues, path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '확인필요'
    hdr = ['구분', '적요', '거래처코드', '거래처명', '원본 품목명', '수량', '오류내용', '검색키']
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        _style_header(ws.cell(row=1, column=c))
    fill = PatternFill('solid', fgColor=HILITE_COLOR)
    for i, x in enumerate(issues, start=2):
        ws.append([x['mode'], x['remark'], x['code'], x['partner'], x['item'], x['qty'], x['error'], x['key']])
        for c in range(1, 9):
            ws.cell(row=i, column=c).fill = fill
        ws.cell(row=i, column=3).number_format = '@'
    for col, w in zip('ABCDEFGH', (8, 30, 12, 16, 36, 10, 40, 36)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(path)


def _style_header(cell):
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=HEADER_COLOR)
    cell.alignment = Alignment(horizontal='center')


# ═══════════ ② 원장 비교 ═══════════
def read_purchase_workbook(path):
    products, display, excluded, shipping, exclusions = {}, {}, {}, {'다모아': 0.0, '제이제이': 0.0, '제이무역': 0.0}, []
    for sheet, rows in _rows(path).items():
        owner = purchase_owner(sheet)
        for r in range(1, len(rows) + 1):
            item = _s(_cell(rows, r, 4)); qty = _num(_cell(rows, r, 7))
            amount = _num(_cell(rows, r, 18)); label = _s(_cell(rows, r, 19))
            if '배송비합계' in label:
                so = shipping_ledger_owner(sheet, owner)
                if so != '제외':
                    shipping[so] += amount
                else:
                    exclusions.append(('배송비', sheet, '용차 배송비', '', amount, '용차 배송비 비교 제외'))
            if item == '':
                continue
            if '파손/재발송' in item:
                base, ex = item, 0.0
                m = re.match(r'^(.*?)\*(\d+(?:\.\d+)?)\s*파손/재발송', item)
                if m:
                    base, ex = m.group(1).strip(), float(m.group(2))
                key = '%s|%s' % (owner, canonical(base, owner))
                excluded[key] = excluded.get(key, 0.0) + ex
                exclusions.append(('품목', sheet, item, ex, '', '파손/재발송 제외'))
                continue
            if item.startswith('천막추가옵션/'):
                exclusions.append(('품목', sheet, item, qty, '', '배송 부가옵션으로 품목 비교 제외')); continue
            if owner == '미분류':
                exclusions.append(('품목', sheet, item, qty, '', '원장 연결 미분류')); continue
            conv = owner_item_name(item, owner)
            key = '%s|%s' % (owner, canonical(conv, owner))
            products[key] = products.get(key, 0.0) + qty
            display.setdefault(key, conv)
    return {'products': products, 'display': display, 'excluded': excluded, 'shipping': shipping, 'exclusions': exclusions}


def _header_row(rows):
    for r in range(1, min(20, len(rows)) + 1):
        if re.search(r'품\s*목\s*명', _s(_cell(rows, r, 4))):
            return r
    return 0


def read_ledger_workbook(path, owner):
    rows = next(iter(_rows(path).values()))
    hr = _header_row(rows)
    if hr == 0:
        raise ValueError("%s 원장에서 '품목명' 열 제목을 찾지 못했습니다." % owner)
    products, display, prices, supply, rows_by = {}, {}, {}, {}, {}
    service = {'Delivery': 0.0, 'Packing': 0.0, 'ThreePL': 0.0}
    service_rows = {'Delivery': [], 'Packing': [], 'ThreePL': []}
    for r in range(hr + 1, len(rows) + 1):
        item = _s(_cell(rows, r, 4))
        if item == '':
            continue
        qty, unit, sup = _num(_cell(rows, r, 6)), _num(_cell(rows, r, 8)), _num(_cell(rows, r, 9))
        low = item.lower()
        if '용차' in low:
            continue
        if '3pl' in low:
            service['ThreePL'] += sup; service_rows['ThreePL'].append(r); continue
        if '포장비' in low:
            service['Packing'] += sup; service_rows['Packing'].append(r); continue
        if '택배비' in low:
            service['Delivery'] += sup; service_rows['Delivery'].append(r); continue
        conv = owner_item_name(item, owner)
        key = '%s|%s' % (owner, canonical(conv, owner))
        products[key] = products.get(key, 0.0) + qty
        supply[key] = supply.get(key, 0.0) + sup
        rows_by.setdefault(key, []).append(r)
        display.setdefault(key, conv)
        prices.setdefault(key, [])
        if unit not in prices[key]:
            prices[key].append(unit)
    return {'rows': rows, 'header_row': hr, 'products': products, 'display': display, 'prices': prices, 'supply': supply,
            'service': service, 'rows_by': rows_by, 'service_rows': service_rows}


def read_order_shipping_categories(order_path):
    rows = _order_rows(order_path)
    groups = {}
    for r in range(2, len(rows) + 1):
        alias = _s(_cell(rows, r, 2))
        if alias == '':
            continue
        carrier = _s(_cell(rows, r, 12)) or '택배사미입력'
        owner, cat = order_shipping_category(alias)
        ship = _num(_cell(rows, r, 18))
        if owner == '다모아':
            ship += _num(_cell(rows, r, 19)) + _num(_cell(rows, r, 20))
        g = groups.setdefault((owner, cat, carrier), {'owner': owner, 'category': cat, 'carrier': carrier, 'count': 0, 'shipping': 0.0})
        g['count'] += 1; g['shipping'] += ship
    return sorted(groups.values(), key=lambda g: (g['owner'], g['category'], g['carrier']))


def build_ledger_comparison(purchase_path, sales_path, order_path, erp_path, ledgers, out_dir=None, log=print):
    """ledgers: {'다모아': path|None, '제이제이': …, '제이무역': …}. 반환 dict(out, rows(품목비교), summary)."""
    for p in (purchase_path, sales_path, order_path, erp_path):
        if not os.path.exists(p):
            raise FileNotFoundError('필요한 파일을 찾을 수 없습니다: ' + p)
    active = []
    for owner in ('다모아', '제이제이', '제이무역'):
        p = (ledgers.get(owner) or '').strip()
        if not p:
            log('%s 원장 공란: 비교 제외' % owner); continue
        if not os.path.exists(p):
            raise FileNotFoundError('%s 원장 파일을 찾을 수 없습니다: %s' % (owner, p))
        active.append(owner)
    out_dir = out_dir or os.path.dirname(purchase_path)
    erp = load_erp_map(erp_path)
    purchase = read_purchase_workbook(purchase_path)
    categories = read_order_shipping_categories(order_path)
    log('주문서 택배비 카테고리 %d개 집계' % len(categories))
    led = {o: read_ledger_workbook(ledgers[o], o) for o in active}

    results, summary = [], {}
    for owner in active:
        summary[owner] = {'total': 0, 'normal': 0, 'mismatch': 0, 'erp_missing': 0}
        keys = {k for k in purchase['products'] if k.startswith(owner + '|')} | set(led[owner]['products'])
        for key in keys:
            summary[owner]['total'] += 1
            p_qty = purchase['products'].get(key, 0.0)
            l_qty = led[owner]['products'].get(key, 0.0)
            ex = purchase['excluded'].get(key, 0.0)
            adj_qty, adj_sup = l_qty, led[owner]['supply'].get(key, 0.0)
            lprices = led[owner]['prices'].get(key, [])
            l_unit = lprices[0] if len(lprices) == 1 else None
            notes = []
            if ex > 0 and _near(l_qty, p_qty + ex):
                adj_qty = l_qty - ex
                if l_unit is not None:
                    adj_sup -= l_unit * ex
                notes.append('파손/재발송 %s 제외' % _s(ex))
            canon = key.split('|', 1)[1]
            er = resolve_erp_item(erp, canon)
            erp_price = er['price']
            expected = p_qty * erp_price if erp_price is not None else None
            diff = adj_sup - expected if expected is not None else None
            status = '정상'
            if p_qty == 0 and adj_qty != 0:
                status = '원장에만 있음'
            elif p_qty != 0 and adj_qty == 0:
                status = '매입에만 있음'
            elif not _near(p_qty, adj_qty):
                status = '수량 불일치'
            if er['status'] != '정상':
                status = er['status']; summary[owner]['erp_missing'] += 1
                if er['note']:
                    notes.append('ERP 단가 후보: ' + er['note'])
            elif len(lprices) != 1:
                status = '원장 단가 중복'; notes.append('원장 단가: ' + ', '.join(_s(x) for x in lprices))
            elif not _near(erp_price, l_unit):
                status = '단가 불일치'
            elif diff is not None and not _near(diff, 0):
                status = '금액 불일치'
            summary[owner]['normal' if status == '정상' else 'mismatch'] += 1
            results.append({'key': key, 'owner': owner, 'status': status,
                            'pname': purchase['display'].get(key, ''), 'lname': led[owner]['display'].get(key, ''),
                            'ename': er['item']['name'] if er['item'] else '', 'pqty': p_qty, 'lqty': l_qty, 'ex': ex,
                            'adjqty': adj_qty, 'eprice': erp_price, 'lprice': l_unit, 'expected': expected,
                            'lsup': adj_sup, 'diff': diff, 'note': ' / '.join(notes)})
    ship_status = {}
    for owner in active:
        sv = led[owner]['service']
        lt = sv['Delivery'] + sv['Packing'] + sv['ThreePL']
        ship_status[owner] = _near(lt - purchase['shipping'][owner], 0)

    # ── 결과 통합문서 ──
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = '품목비교'
    hdr = ['원장', '상태', '매입 품목', '원장 품목', 'ERP 품목', '매입수량', '원장수량', '제외수량', '비교 원장수량', 'ERP 구매단가', '원장단가', '예상 공급가액', '원장 공급가액', '금액차이', '비고']
    ws.append(hdr)
    for c in range(1, 16):
        _style_header(ws.cell(row=1, column=c))
    results.sort(key=lambda x: (x['owner'], x['status'], x['pname'], x['lname']))
    mism = PatternFill('solid', fgColor=MISMATCH_COLOR)
    for i, x in enumerate(results, start=2):
        ws.append([x['owner'], x['status'], x['pname'], x['lname'], x['ename'], x['pqty'], x['lqty'], x['ex'], x['adjqty'],
                   x['eprice'], x['lprice'], x['expected'], x['lsup'], x['diff'], x['note']])
        fill = PatternFill('solid', fgColor=OWNER_COLORS.get(x['owner'], 'F2F2F2'))
        for c in range(1, 16):
            ws.cell(row=i, column=c).fill = fill
            if 6 <= c <= 14:
                ws.cell(row=i, column=c).number_format = '#,##0.###'
        if x['status'] != '정상':
            ws.cell(row=i, column=2).fill = mism; ws.cell(row=i, column=2).font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJKLMNO', (9, 14, 34, 34, 34, 10, 10, 10, 12, 12, 10, 14, 14, 12, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

    # 매입매출자료
    ws2 = wb.create_sheet('매입매출자료')
    ws2.append(['시트명', '구분', '', 'R열 금액', 'S열 내용'])
    for c in range(1, 6):
        _style_header(ws2.cell(row=1, column=c))
    r = 2
    for path, kind, colr in ((purchase_path, '매입', PURCHASE_SUMMARY_COLOR), (sales_path, '매출', SALES_SUMMARY_COLOR)):
        for sheet, rows in _rows(path).items():
            for rr in range(1, len(rows) + 1):
                amount, label = _cell(rows, rr, 18), _s(_cell(rows, rr, 19))
                if (amount is None or _s(amount) == '') and label == '':
                    continue
                ws2.append([sheet, kind, '', amount if _s(amount) != '' else None, label])
                for c in range(1, 6):
                    ws2.cell(row=r, column=c).fill = PatternFill('solid', fgColor=colr)
                ws2.cell(row=r, column=4).number_format = '#,##0.###'
                r += 1
    for col, w in zip('ABCDE', (28, 10, 3, 18, 38)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = 'A2'

    # 택배비카테고리
    ws3 = wb.create_sheet('택배비카테고리')
    ws3.append(['거래처', '카테고리-택배사', '카테고리', '택배사', '주문건수', '배송비 합계'])
    for c in range(1, 7):
        _style_header(ws3.cell(row=1, column=c))
    for i, g in enumerate(categories, start=2):
        ws3.append([g['owner'], '%s-%s' % (g['category'], g['carrier']), g['category'], g['carrier'], g['count'], g['shipping']])
        for c in range(1, 7):
            ws3.cell(row=i, column=c).fill = PatternFill('solid', fgColor=OWNER_COLORS.get(g['owner'], 'F2F2F2'))
        ws3.cell(row=i, column=5).number_format = '#,##0'; ws3.cell(row=i, column=6).number_format = '#,##0'
    for col, w in zip('ABCDEF', (12, 30, 24, 14, 12, 16)):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = 'A2'

    # 원장-* (값 복사 + 불일치 행 강조 + O열 요청합계 + Q:S 택배비 카테고리)
    for owner in active:
        hi = set()
        for x in results:
            if x['owner'] == owner and x['status'] != '정상':
                hi.update(led[owner]['rows_by'].get(x['key'], []))
        if not ship_status[owner]:
            for k in ('Delivery', 'Packing', 'ThreePL'):
                hi.update(led[owner]['service_rows'][k])
        _write_ledger_sheet(wb.create_sheet('원장-' + owner), led[owner], owner, hi, categories)

    # 주문서-* (대량등록양식·품목별토탈수량·주문서 붙여넣기 제외 → 보통 없음)
    for sheet, rows in _rows(order_path).items():
        if sheet in ('대량등록양식', '품목별토탈수량', '주문서 붙여넣기', '주문서붙여넣기'):
            continue
        wsx = wb.create_sheet(('주문서-' + sheet)[:31])
        for rr in rows:
            wsx.append(list(rr))
    out_path = os.path.join(out_dir, '원장비교결과_%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    wb.save(out_path)
    log('원장 비교 완료: %s' % os.path.basename(out_path))
    return {'out': out_path, 'rows': results, 'summary': summary, 'excluded_owners': [o for o in ('다모아', '제이제이', '제이무역') if o not in active]}


def _write_ledger_sheet(ws, ledger, owner, highlight_rows, categories):
    rows, hr = ledger['rows'], ledger['header_row']
    for rr in rows:
        ws.append(list(rr))
    hl = PatternFill('solid', fgColor=HILITE_COLOR)
    last_col = max(15, max((len(rr) for rr in rows), default=1))
    own_cats = [g for g in categories if g['owner'] == owner]
    if own_cats:
        last_col = max(19, last_col)
    for r in sorted(highlight_rows):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).fill = hl
    # 요청 합계 (O열)
    if owner in ('다모아', '제이제이') and hr:
        general = canopy = delivery = 0.0; g_row = c_row = d_row = 0
        for r in range(hr + 1, len(rows) + 1):
            item, note, sup = _s(_cell(rows, r, 4)), _s(_cell(rows, r, 14)), _num(_cell(rows, r, 9))
            if owner == '다모아':
                if not re.match(r'^(통합택배비|포장비|3\s*PL\s*/?\s*수작업비용\s*\(다모아\)|수작업비용\s*\(다모아\))$', item):
                    continue
                if re.search(r'천막\s*(택배|포장)', note):
                    canopy += sup; c_row = c_row or r
                elif note == '':
                    general += sup; g_row = g_row or r
            elif re.match(r'^(택배비|화물택배비)$', item):
                delivery += sup; d_row = d_row or r
        ws.cell(row=hr, column=15, value='요청 합계')
        for rr_, amt in ((g_row, general), (c_row, canopy), (d_row, delivery)):
            if rr_:
                cell = ws.cell(row=rr_, column=15, value=amt)
                cell.number_format = '#,##0'; cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='E2EFDA')
        ws.column_dimensions['O'].width = 14
    # 주문서 택배비 카테고리 표 (Q:S)
    if own_cats:
        hr_ = hr or 2
        head_color = {'다모아': '5B9BD5', '제이제이': '70AD47', '제이무역': 'ED7D31'}[owner]
        row_color = OWNER_COLORS[owner]
        t = ws.cell(row=max(1, hr_ - 1), column=17, value='주문서 택배비 카테고리')
        ws.merge_cells(start_row=max(1, hr_ - 1), start_column=17, end_row=max(1, hr_ - 1), end_column=19)
        t.font = Font(bold=True, color='FFFFFF'); t.fill = PatternFill('solid', fgColor=head_color); t.alignment = Alignment(horizontal='center')
        for c, v in zip((17, 18, 19), ('카테고리-택배사', '주문건수', '배송비 합계')):
            cell = ws.cell(row=hr_, column=c, value=v); cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor=row_color)
            cell.alignment = Alignment(horizontal='center')
        r, total = hr_ + 1, 0.0
        for g in sorted(own_cats, key=lambda g: (g['category'], g['carrier'])):
            for c, v in zip((17, 18, 19), ('%s-%s' % (g['category'], g['carrier']), g['count'], g['shipping'])):
                cell = ws.cell(row=r, column=c, value=v); cell.fill = PatternFill('solid', fgColor=row_color)
            total += g['shipping']; r += 1
        ws.cell(row=r, column=17, value='합계'); ws.cell(row=r, column=19, value=total)
        for c in (17, 18, 19):
            ws.cell(row=r, column=c).font = Font(bold=True); ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=head_color)
        ws.column_dimensions['Q'].width = 30; ws.column_dimensions['R'].width = 11; ws.column_dimensions['S'].width = 16
        for rr_ in range(hr_ + 1, r + 1):
            ws.cell(row=rr_, column=18).number_format = '#,##0'; ws.cell(row=rr_, column=19).number_format = '#,##0'


# ═══════════ ③ 제이 매입 품목 합산 ═══════════
def build_jay_purchase_summary(purchase_path, template_path, trade_date, out_dir=None, log=print):
    if not os.path.exists(template_path):
        raise FileNotFoundError('제이 품목거래 양식파일을 찾을 수 없습니다: ' + template_path)
    items, order, matched, src_rows = {}, [], [], 0
    for sheet, rows in _rows(purchase_path).items():
        n = _s(sheet)
        if not (n in ('제이스토어', '제이쿠팡') or re.match(r'^제이개인(?:$|[\s/_-])', n)):
            continue
        matched.append(n)
        for r in range(1, len(rows) + 1):
            name = _s(_cell(rows, r, 4))
            if name == '':
                continue
            src_rows += 1
            loc, qty = _s(_cell(rows, r, 1)), _num(_cell(rows, r, 7))
            if name not in items:
                items[name] = {'loc': loc, 'qty': qty}; order.append(name)
            else:
                items[name]['qty'] += qty
                if not items[name]['loc'] and loc:
                    items[name]['loc'] = loc
    if not matched:
        raise ValueError("매입파일에서 '제이스토어', '제이쿠팡', '제이개인' 시트를 찾지 못했습니다.")
    if not order:
        raise ValueError('대상 시트의 D열에서 합산할 품목을 찾지 못했습니다.')
    out_dir = out_dir or os.path.dirname(purchase_path)
    out_path = os.path.join(out_dir, '제이 품목거래 %d-%d.xls' % (trade_date.month, trade_date.day))
    date_txt = '%d월 %d일' % (trade_date.month, trade_date.day)
    rb = xlrd.open_workbook(template_path, formatting_info=True)
    wb = xl_copy(rb)
    ws = wb.get_sheet(rb.sheet_names().index('대량등록양식'))
    num = _xls_style(numfmt='#,##0'); plain = _xls_style()
    for i, name in enumerate(order):
        ws.write(i, 0, '제이제이컴퍼니(유)', plain)
        ws.write(i, 3, name, plain)
        ws.write(i, 6, float(_round(items[name]['qty'])), num)
        ws.write(i, 17, date_txt, plain)
    wb.save(out_path)
    log('제이 매입 합산 완료: 대상 %d개 시트, 원본 %d행 → 품목 %d개' % (len(matched), src_rows, len(order)))
    return {'out': out_path, 'items': [(n, items[n]['qty']) for n in order], 'sheets': matched}


# ═══════════ 폴더 자동 탐색 (도구의 '폴더에서 자동 불러오기') ═══════════
def find_files_in_folder(folder):
    files = []
    for f in os.listdir(folder):
        p = os.path.join(folder, f)
        base, ext = os.path.splitext(f)
        if not os.path.isfile(p) or ext.lower() not in ('.xls', '.xlsx', '.xlsm') or f.startswith('~$'):
            continue
        if re.search(r'원장비교결과|ERP_판매구매일괄등록|ERP_확인필요|시트분리|품목정보관리|판매구매일괄등록|제이 품목거래', base):
            continue
        files.append((base, p, os.path.getmtime(p)))

    def latest(inc, exc=''):
        hits = [x for x in files if re.search(inc, x[0]) and not (exc and re.search(exc, x[0]))]
        return max(hits, key=lambda x: x[2])[1] if hits else ''
    return {
        'purchase': latest(r'(^|[\s_.-])매입|^매입', '매입매출'),
        'sales': latest(r'(^|[\s_.-])매출|^매출', '매입매출'),
        'damoa': latest(r'(판매처.*원장|다모아.*원장)', '제이무역|무역'),
        'jj': latest(r'(제이제이.*원장|^청년몰[\s_.-]+\d)', '청년몰원장|제이무역|무역'),
        'trade': latest(r'(청년몰원장|제이무역.*원장|무역.*원장)'),
        'order': latest(r'주문서', '원장비교결과|시트분리'),
    }
